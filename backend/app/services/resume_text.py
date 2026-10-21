"""Upload validation (by content, never by client type/extension) and text extraction."""
import io
import logging
import zipfile
from dataclasses import dataclass

from pypdf import PdfReader

from app.exceptions import AppError, UnprocessableError

logger = logging.getLogger(__name__)

MAX_DECOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
MAX_PDF_PAGES = 100
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class PayloadTooLargeError(AppError):
    """413."""

    status_code = 413


class UnsupportedTypeError(AppError):
    """415."""

    status_code = 415


@dataclass(frozen=True)
class ParsedDocument:
    """Result of validating and extracting a file."""

    content_type: str
    text: str
    page_count: int | None


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _check_zip(data: bytes) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        infos = archive.infolist()
    except (zipfile.BadZipFile, OSError) as exc:
        raise UnprocessableError("The DOCX file is corrupt") from exc
    if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_DECOMPRESSED_BYTES:
        raise UnprocessableError("The document expands to an unsafe size and was rejected")
    return archive


def _parse_pdf(data: bytes) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnprocessableError("Password-protected PDFs are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise UnprocessableError(f"PDF has more than {MAX_PDF_PAGES} pages")
        pages = [page.extract_text() or "" for page in reader.pages]
    except AppError:
        raise
    except Exception as exc:  # pypdf raises many error types for malformed input
        raise UnprocessableError("The PDF file is corrupt or unreadable") from exc
    return ParsedDocument(PDF_MIME, _normalise("\n".join(pages)), len(pages))


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx  # local import keeps module import cheap

    _check_zip(data)
    try:
        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    except Exception as exc:
        raise UnprocessableError("The DOCX file is corrupt or unreadable") from exc
    return ParsedDocument(DOCX_MIME, _normalise("\n".join(lines)), None)


def parse_upload(data: bytes) -> ParsedDocument:
    """Detect type via magic bytes/structure, then extract text.

    Raises UnsupportedTypeError (415) or UnprocessableError (422).
    """
    if data.startswith(b"%PDF-"):
        return _parse_pdf(data)
    if data.startswith(b"PK") and zipfile.is_zipfile(io.BytesIO(data)):
        archive = _check_zip(data)
        if "word/document.xml" in archive.namelist():
            return _parse_docx(data)
    raise UnsupportedTypeError("Only PDF and DOCX files are accepted")
